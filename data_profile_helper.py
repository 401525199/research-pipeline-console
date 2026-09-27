# SPDX-License-Identifier: GPL-3.0-only
# Copyright (C) 2026 Qi Wang

from __future__ import annotations

import csv
import json
import math
import sys
import warnings
from pathlib import Path
from typing import Any

import pandas as pd


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


MAX_PROFILE_ROWS = 50_000
MAX_SAMPLE_ROWS = 12
MAX_SAMPLE_COLUMNS = 12
MAX_FIELD_ROWS = 120


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    if isinstance(value, (bool, int, float, str)):
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value
    return str(value)


def _short(value: Any, limit: int = 120) -> str:
    text = " ".join(str(_json_value(value) or "").split())
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _semantic_type(series: pd.Series) -> str:
    if pd.api.types.is_bool_dtype(series.dtype):
        return "布尔"
    if pd.api.types.is_integer_dtype(series.dtype):
        return "整数"
    if pd.api.types.is_float_dtype(series.dtype):
        return "数值"
    if pd.api.types.is_datetime64_any_dtype(series.dtype):
        return "日期时间"
    if isinstance(series.dtype, pd.CategoricalDtype):
        return "分类"
    sample = series.dropna().astype(str).head(300)
    if sample.empty:
        return "空列 / 未识别"
    numeric = pd.to_numeric(sample.str.replace(",", "", regex=False), errors="coerce")
    if float(numeric.notna().mean()) >= 0.95:
        return "数值（文本存储）"
    name = str(series.name).casefold()
    if any(token in name for token in ("date", "time", "month", "quarter")):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            parsed = pd.to_datetime(sample, errors="coerce")
        if float(parsed.notna().mean()) >= 0.80:
            return "日期（文本存储）"
    return "文本"


def _candidate_key(frame: pd.DataFrame) -> tuple[list[str], int | None, str]:
    lookup = {str(column).strip().casefold(): str(column) for column in frame.columns}
    single_candidates = (
        "entity_id",
        "subject_id",
        "case_id",
        "row_id",
        "observation_id",
        "record_id",
        "id",
    )
    pair_candidates = (
        ("entity_id", "year"),
        ("entity_id", "date"),
        ("subject_id", "year"),
        ("subject_id", "date"),
        ("case_id", "period"),
        ("record_id", "period"),
        ("code", "year"),
        ("code", "date"),
    )
    candidates: list[list[str]] = []
    for name in single_candidates:
        if name in lookup:
            candidates.append([lookup[name]])
    for names in pair_candidates:
        if all(name in lookup for name in names):
            candidates.append([lookup[name] for name in names])
    if not candidates:
        return [], None, "未识别候选主键"

    best_columns: list[str] = []
    best_duplicates: int | None = None
    best_score: tuple[int, int] | None = None
    for columns in candidates:
        values = frame[columns]
        complete = values.notna().all(axis=1)
        complete_values = values.loc[complete]
        duplicates = int(complete_values.duplicated(keep=False).sum()) if not complete_values.empty else 0
        missing = int((~complete).sum())
        score = (duplicates, missing)
        if best_score is None or score < best_score:
            best_score = score
            best_columns = columns
            best_duplicates = duplicates
    scope = "已扫描样本内" if len(frame) >= MAX_PROFILE_ROWS else "当前文件内"
    return best_columns, best_duplicates, scope


def _range_summary(frame: pd.DataFrame) -> tuple[str, str]:
    year_summary = "未识别年份字段"
    date_summary = "未识别日期字段"
    for column in frame.columns:
        name = str(column).casefold()
        series = frame[column]
        if year_summary == "未识别年份字段" and any(token in name for token in ("year", "fyear")):
            values = pd.to_numeric(series, errors="coerce").dropna()
            values = values[(values >= 1800) & (values <= 2200)]
            if not values.empty:
                year_summary = f"{str(column)}：{int(values.min())}–{int(values.max())}"
        if date_summary == "未识别日期字段" and any(token in name for token in ("date", "time", "month")):
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                values = pd.to_datetime(series, errors="coerce").dropna()
            if not values.empty:
                date_summary = f"{str(column)}：{values.min().date()}–{values.max().date()}"
    return year_summary, date_summary


def _profile_frame(
    frame: pd.DataFrame,
    *,
    total_rows: int | None,
    source_detail: str,
    sheets: list[str] | None = None,
    labels: dict[str, str] | None = None,
) -> dict[str, Any]:
    frame = frame.copy()
    frame.columns = [str(column).strip() or f"column_{index + 1}" for index, column in enumerate(frame.columns)]
    scanned_rows = int(len(frame))
    sampled = bool(total_rows is None or total_rows > scanned_rows)
    total_cells = max(1, scanned_rows * max(1, len(frame.columns)))
    total_missing = int(frame.isna().sum().sum()) if scanned_rows else 0
    overall_missing_rate = total_missing / total_cells if scanned_rows else 0.0
    key_columns, duplicate_count, duplicate_scope = _candidate_key(frame)
    year_range, date_range = _range_summary(frame)
    labels = labels or {}

    fields: list[dict[str, Any]] = []
    for column in list(frame.columns)[:MAX_FIELD_ROWS]:
        series = frame[column]
        missing_count = int(series.isna().sum())
        fields.append(
            {
                "name": column,
                "dtype": _semantic_type(series),
                "missing_count": missing_count,
                "missing_rate": missing_count / max(1, scanned_rows),
                "unique_count": int(series.nunique(dropna=True)),
                "label": _short(labels.get(column, ""), 80),
            }
        )

    preferred: list[str] = []
    for column in key_columns:
        if column not in preferred:
            preferred.append(column)
    for column in frame.columns:
        lowered = str(column).casefold()
        if any(token in lowered for token in ("year", "date", "value", "score", "label", "category")) and column not in preferred:
            preferred.append(column)
    for column in frame.columns:
        if column not in preferred:
            preferred.append(column)
    sample_columns = preferred[:MAX_SAMPLE_COLUMNS]
    sample_rows = [
        [_short(value) for value in row]
        for row in frame[sample_columns].head(MAX_SAMPLE_ROWS).itertuples(index=False, name=None)
    ] if sample_columns else []

    return {
        "status": "ok",
        "source_detail": source_detail,
        "rows_total": total_rows,
        "rows_scanned": scanned_rows,
        "sampled": sampled,
        "columns_total": int(len(frame.columns)),
        "overall_missing_rate": overall_missing_rate,
        "candidate_key": key_columns,
        "duplicate_count": duplicate_count,
        "duplicate_scope": duplicate_scope,
        "year_range": year_range,
        "date_range": date_range,
        "sheets": sheets or [],
        "fields": fields,
        "sample_columns": sample_columns,
        "sample_rows": sample_rows,
    }


def _csv_settings(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()[:131072]
    text = ""
    encoding = "utf-8-sig"
    for candidate in ("utf-8-sig", "utf-8", "gb18030", "cp1252"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    delimiter = "\t" if path.suffix.casefold() == ".tsv" else ","
    try:
        delimiter = csv.Sniffer().sniff(text, delimiters=",\t;|").delimiter
    except csv.Error:
        pass
    return encoding, delimiter


def _profile_csv(path: Path) -> dict[str, Any]:
    encoding, delimiter = _csv_settings(path)
    kwargs = {
        "sep": delimiter,
        "encoding": encoding,
        "on_bad_lines": "skip",
        "low_memory": False,
    }
    frame = pd.read_csv(path, nrows=MAX_PROFILE_ROWS, **kwargs)
    total_rows: int | None = None
    if frame.columns.size:
        try:
            total_rows = 0
            for chunk in pd.read_csv(
                path,
                usecols=[0],
                chunksize=200_000,
                sep=delimiter,
                encoding=encoding,
                on_bad_lines="skip",
            ):
                total_rows += len(chunk)
        except (OSError, ValueError, pd.errors.ParserError):
            total_rows = None
    delimiter_name = {",": "逗号", "\t": "制表符", ";": "分号", "|": "竖线"}.get(delimiter, repr(delimiter))
    return _profile_frame(
        frame,
        total_rows=total_rows,
        source_detail=f"{delimiter_name}分隔 · {encoding}",
    )


def _profile_excel(path: Path) -> dict[str, Any]:
    if path.suffix.casefold() == ".xls":
        raise RuntimeError("旧式 XLS 需要额外读取组件；请另存为 XLSX 后进行深度画像")
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    sheets = [str(name) for name in book.sheetnames]
    if not sheets:
        raise RuntimeError("工作簿没有可读取的工作表")
    sheet_sizes = {
        name: (int(book[name].max_row or 0), int(book[name].max_column or 0))
        for name in sheets
    }
    data_sheet = max(sheets, key=lambda name: sheet_sizes[name][0] * max(1, sheet_sizes[name][1]))
    max_row = sheet_sizes[data_sheet][0]
    book.close()
    preview = pd.read_excel(path, sheet_name=data_sheet, header=None, nrows=12, engine="openpyxl")
    header_row = int(preview.notna().sum(axis=1).idxmax()) if not preview.empty else 0
    frame = pd.read_excel(
        path,
        sheet_name=data_sheet,
        header=header_row,
        nrows=MAX_PROFILE_ROWS,
        engine="openpyxl",
    )
    total_rows: int | None = max(0, max_row - header_row - 1)
    return _profile_frame(
        frame,
        total_rows=total_rows,
        source_detail=f"主要工作表：{data_sheet} · 表头行：{header_row + 1}",
        sheets=sheets,
    )


def _profile_stata(path: Path) -> dict[str, Any]:
    metadata = pd.io.stata.StataReader(path)
    labels = {str(key): str(value) for key, value in metadata.variable_labels().items()}
    data_label = str(metadata.data_label or "").strip()
    frame: pd.DataFrame | None = None
    total_rows = 0
    for chunk in pd.read_stata(path, iterator=True, chunksize=MAX_PROFILE_ROWS):
        total_rows += len(chunk)
        if frame is None:
            frame = chunk
    if frame is None:
        frame = pd.DataFrame(columns=list(labels))
    detail = f"Stata 数据集{f' · {data_label}' if data_label else ''}"
    return _profile_frame(frame, total_rows=total_rows, source_detail=detail, labels=labels)


def profile(path: Path) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        return _profile_csv(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return _profile_excel(path)
    if suffix == ".dta":
        return _profile_stata(path)
    if suffix == ".sav":
        frame = pd.read_spss(path).head(MAX_PROFILE_ROWS)
        return _profile_frame(frame, total_rows=None, source_detail="SPSS 数据集")
    if suffix == ".parquet":
        frame = pd.read_parquet(path).head(MAX_PROFILE_ROWS)
        return _profile_frame(frame, total_rows=None, source_detail="Parquet 数据集")
    if suffix == ".feather":
        frame = pd.read_feather(path).head(MAX_PROFILE_ROWS)
        return _profile_frame(frame, total_rows=None, source_detail="Feather 数据集")
    if suffix == ".pkl":
        return {
            "status": "limited",
            "message": "为避免执行不可信的序列化对象，PKL 文件不自动加载。",
        }
    return {"status": "limited", "message": f"暂不支持 {suffix or '无扩展名'} 的深度数据画像。"}


def main() -> int:
    if len(sys.argv) != 2:
        print(json.dumps({"status": "error", "message": "缺少数据文件路径"}, ensure_ascii=False))
        return 2
    path = Path(sys.argv[1])
    try:
        result = profile(path)
    except Exception as exc:  # Return a concise UI-safe error instead of a traceback.
        result = {"status": "error", "message": f"{type(exc).__name__}: {exc}"}
    print(json.dumps(result, ensure_ascii=False, default=_json_value))
    return 0 if result.get("status") != "error" else 1


if __name__ == "__main__":
    raise SystemExit(main())
